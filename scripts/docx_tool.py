#!/usr/bin/env python3
"""
docx_tool.py — Unified tool for creating, reading, editing, and converting .docx documents.

This is a self-contained script.  No imports from other files are required.

Python API::

    from docx_tool import DocxBuilder, DocxReader, DocxEditor, DocxUtil

    # Create
    builder = DocxBuilder().preset_chinese().add_heading("标题").add_paragraph("正文").save("out.docx")

    # Read
    reader = DocxReader("report.docx")
    reader.extract_to_markdown("report.md")
    reader.print_structure()
    reader.grep("关键词", context=2)
    reader.extract_range(10, 25, "section.md")
    reader.extract_section("Results", "results.md")

    # Edit
    editor = DocxEditor("in.docx")
    editor.replace("old", "new")
    editor.insert_after(3, "新段落")
    editor.delete(5)
    editor.set_style(2, "Heading 1")
    editor.save("out.docx")

    # Utils
    DocxUtil.merge_runs("in.docx", "out.docx")
    DocxUtil.convert_to_doc("in.docx", "out.doc")
    DocxUtil.apply_theme("in.docx", "theme.yaml", "out.docx")

CLI::

    # Create
    python docx_tool.py create --output out.docx --preset chinese --title "报告" --add-heading "1:概述"

    # Read/Extract
    python docx_tool.py extract report.docx --output report.md
    python docx_tool.py extract report.docx --structure
    python docx_tool.py extract report.docx --grep "关键词" --context 2
    python docx_tool.py extract report.docx --range 10-25 --output section.md
    python docx_tool.py extract report.docx --section "Results"

    # Edit
    python docx_tool.py edit in.docx out.docx --replace "old" "new"
    python docx_tool.py edit in.docx out.docx --insert-after 3 "新段落"
    python docx_tool.py edit in.docx out.docx --delete 5
    python docx_tool.py edit in.docx out.docx --set-style 2 "Heading 1"

    # Utils
    python docx_tool.py merge-runs in.docx out.docx
    python docx_tool.py convert-to-doc in.docx out.doc
    python docx_tool.py apply-theme in.docx theme.yaml out.docx
"""
from __future__ import annotations

import argparse
import copy
import io
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from lxml import etree  # type: ignore[attr-defined]

# ═══════════════════════════════════════════════════════════════
# Inline LaTeX → OMML (from latex2omml.py)
# ═══════════════════════════════════════════════════════════════

try:
    from latex2mathml import converter as _latex_converter  # type: ignore[import]
except ImportError:
    _latex_converter = None

_OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_MATHML_NS = "http://www.w3.org/1998/Math/MathML"
_XML_NS = "http://www.w3.org/XML/1998/namespace"
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _omml_tag(elem):
    return elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag


def _omml_extract_text(elem):
    parts = []
    if elem.text:
        parts.append(elem.text)
    for c in elem:
        parts.append(_omml_extract_text(c))
        if c.tail:
            parts.append(c.tail)
    return "".join(parts)


def _omml_add_text_run(parent, text):
    from xml.etree.ElementTree import SubElement
    r = SubElement(parent, f"{{{_OMML_NS}}}r")
    t = SubElement(r, f"{{{_OMML_NS}}}t")
    t.text = text
    if text.startswith(" ") or text.endswith(" "):
        t.set(f"{{{_XML_NS}}}space", "preserve")


def _omml_convert(ml_elem, omml_parent):
    from xml.etree.ElementTree import SubElement
    tag = _omml_tag(ml_elem)
    kids = list(ml_elem)

    if tag == "mrow":
        for c in kids:
            _omml_convert(c, omml_parent)

    elif tag == "msup":
        e = SubElement(omml_parent, f"{{{_OMML_NS}}}sSup")
        SubElement(e, f"{{{_OMML_NS}}}sSupPr")
        ce = SubElement(e, f"{{{_OMML_NS}}}e")
        cs = SubElement(e, f"{{{_OMML_NS}}}sup")
        if kids:
            _omml_convert(kids[0], ce)
        if len(kids) > 1:
            _omml_convert(kids[1], cs)

    elif tag == "msub":
        e = SubElement(omml_parent, f"{{{_OMML_NS}}}sSub")
        SubElement(e, f"{{{_OMML_NS}}}sSubPr")
        ce = SubElement(e, f"{{{_OMML_NS}}}e")
        cs = SubElement(e, f"{{{_OMML_NS}}}sub")
        if kids:
            _omml_convert(kids[0], ce)
        if len(kids) > 1:
            _omml_convert(kids[1], cs)

    elif tag == "msubsup":
        e = SubElement(omml_parent, f"{{{_OMML_NS}}}sSubSup")
        SubElement(e, f"{{{_OMML_NS}}}sSubSupPr")
        ce = SubElement(e, f"{{{_OMML_NS}}}e")
        csub = SubElement(e, f"{{{_OMML_NS}}}sub")
        csup = SubElement(e, f"{{{_OMML_NS}}}sup")
        if kids:
            _omml_convert(kids[0], ce)
        if len(kids) > 1:
            _omml_convert(kids[1], csub)
        if len(kids) > 2:
            _omml_convert(kids[2], csup)

    elif tag == "mfrac":
        e = SubElement(omml_parent, f"{{{_OMML_NS}}}f")
        SubElement(e, f"{{{_OMML_NS}}}fPr")
        num = SubElement(e, f"{{{_OMML_NS}}}num")
        den = SubElement(e, f"{{{_OMML_NS}}}den")
        if kids:
            _omml_convert(kids[0], num)
        if len(kids) > 1:
            _omml_convert(kids[1], den)

    elif tag == "msqrt":
        e = SubElement(omml_parent, f"{{{_OMML_NS}}}rad")
        SubElement(e, f"{{{_OMML_NS}}}radPr")
        SubElement(e, f"{{{_OMML_NS}}}deg")
        ce = SubElement(e, f"{{{_OMML_NS}}}e")
        for c in kids:
            _omml_convert(c, ce)

    elif tag == "mroot":
        e = SubElement(omml_parent, f"{{{_OMML_NS}}}rad")
        SubElement(e, f"{{{_OMML_NS}}}radPr")
        deg = SubElement(e, f"{{{_OMML_NS}}}deg")
        ce = SubElement(e, f"{{{_OMML_NS}}}e")
        if len(kids) > 1:
            _omml_convert(kids[1], deg)
        if kids:
            _omml_convert(kids[0], ce)

    elif tag == "munderover":
        op = _omml_extract_text(kids[0]) if kids else ""
        nary_ops = {
            "\u222B": "\u222B",
            "\u2211": "\u2211",
            "\u220F": "\u220F",
            "\u222E": "\u222E",
        }
        if op in nary_ops:
            e = SubElement(omml_parent, f"{{{_OMML_NS}}}nary")
            pr = SubElement(e, f"{{{_OMML_NS}}}naryPr")
            SubElement(pr, f"{{{_OMML_NS}}}chr").set(f"{{{_OMML_NS}}}val", nary_ops[op])
            se = SubElement(e, f"{{{_OMML_NS}}}sub")
            pe = SubElement(e, f"{{{_OMML_NS}}}sup")
            ee = SubElement(e, f"{{{_OMML_NS}}}e")
            if len(kids) > 1:
                _omml_convert(kids[1], se)
            if len(kids) > 2:
                _omml_convert(kids[2], pe)
            if len(kids) > 3:
                _omml_convert(kids[3], ee)
        else:
            e = SubElement(omml_parent, f"{{{_OMML_NS}}}sSubSup")
            SubElement(e, f"{{{_OMML_NS}}}sSubSupPr")
            ce = SubElement(e, f"{{{_OMML_NS}}}e")
            csub = SubElement(e, f"{{{_OMML_NS}}}sub")
            csup = SubElement(e, f"{{{_OMML_NS}}}sup")
            if kids:
                _omml_convert(kids[0], ce)
            if len(kids) > 1:
                _omml_convert(kids[1], csub)
            if len(kids) > 2:
                _omml_convert(kids[2], csup)

    elif tag == "munder":
        e = SubElement(omml_parent, f"{{{_OMML_NS}}}sSub")
        SubElement(e, f"{{{_OMML_NS}}}sSubPr")
        ce = SubElement(e, f"{{{_OMML_NS}}}e")
        cs = SubElement(e, f"{{{_OMML_NS}}}sub")
        if kids:
            _omml_convert(kids[0], ce)
        if len(kids) > 1:
            _omml_convert(kids[1], cs)

    elif tag == "mover":
        e = SubElement(omml_parent, f"{{{_OMML_NS}}}sSup")
        SubElement(e, f"{{{_OMML_NS}}}sSupPr")
        ce = SubElement(e, f"{{{_OMML_NS}}}e")
        cs = SubElement(e, f"{{{_OMML_NS}}}sup")
        if kids:
            _omml_convert(kids[0], ce)
        if len(kids) > 1:
            _omml_convert(kids[1], cs)

    elif tag in ("mi", "mn", "mo", "mtext"):
        t = _omml_extract_text(ml_elem)
        if t:
            _omml_add_text_run(omml_parent, t)

    elif tag == "mspace":
        _omml_add_text_run(omml_parent, " ")

    elif tag == "mfenced":
        op = ml_elem.get("open", "(")
        cl = ml_elem.get("close", ")")
        _omml_add_text_run(omml_parent, op)
        for c in kids:
            _omml_convert(c, omml_parent)
        _omml_add_text_run(omml_parent, cl)

    elif tag == "none":
        pass

    else:
        for c in kids:
            _omml_convert(c, omml_parent)


def _mathml_to_omml_inner(mathml_root, root_elem):
    """Fill root_elem with oMath children from MathML."""
    from xml.etree.ElementTree import SubElement
    for child in list(mathml_root):
        t = _omml_tag(child)
        if t in ("mrow", "mstyle"):
            om = SubElement(root_elem, f"{{{_OMML_NS}}}oMath")
            _omml_convert(child, om)
        elif t == "semantics":
            for sc in child:
                if _omml_tag(sc) == "mrow":
                    om = SubElement(root_elem, f"{{{_OMML_NS}}}oMath")
                    _omml_convert(sc, om)
                    break
    return root_elem


def _mathml_to_omml(mathml_root):
    from xml.etree.ElementTree import Element
    return _mathml_to_omml_inner(mathml_root, Element(f"{{{_OMML_NS}}}oMathPara"))


def _mathml_to_omml_inline(mathml_root):
    from xml.etree.ElementTree import Element
    return _mathml_to_omml_inner(mathml_root, Element(f"{{{_OMML_NS}}}oMath"))


def latex_to_omml(latex_str, inline=False):
    """Convert a LaTeX string to OMML XML.

    Args:
        latex_str: LaTeX source (e.g., ``r"E = mc^2"``).
        inline: If True, return inline ``<m:oMath>``; otherwise block ``<m:oMathPara>``.

    Returns:
        OMML XML string.
    """
    if _latex_converter is None:
        raise ImportError("latex2mathml is required for formula conversion. Install: pip install latex2mathml")
    ml = _latex_converter.convert(latex_str)
    if "xmlns" not in ml[:200]:
        ml = ml.replace("<math", f'<math xmlns="{_MATHML_NS}"', 1)
    from xml.etree.ElementTree import fromstring, tostring
    converter = _mathml_to_omml_inline if inline else _mathml_to_omml
    return tostring(converter(fromstring(ml)), encoding="unicode")


def insert_formula(paragraph, latex_str):
    """Insert an INLINE formula into a paragraph."""
    run = paragraph.runs[-1] if paragraph.runs else paragraph.add_run("")
    omml_xml = latex_to_omml(latex_str, inline=True)
    run._element.append(etree.fromstring(omml_xml))
    return run


def insert_formula_display(paragraph, latex_str):
    """Insert a DISPLAY (block-level) formula into a paragraph."""
    p_elem = paragraph._p
    omml_xml = latex_to_omml(latex_str, inline=False)
    omp_elem = etree.fromstring(omml_xml)
    p_elem.append(omp_elem)
    return omp_elem


# ═══════════════════════════════════════════════════════════════
# Theme helpers (from apply_theme.py)
# ═══════════════════════════════════════════════════════════════

def _theme_size(s):
    s = s.strip()
    if s.endswith("pt"):
        return Pt(float(s[:-2]))
    if s.endswith("cm"):
        return Cm(float(s[:-2]))
    if s.endswith("mm"):
        return Cm(float(s[:-2]) / 10)
    if s.endswith("in"):
        return Inches(float(s[:-2]))
    return s


def _theme_get(theme, section, key, default=None):
    return theme.get(section, {}).get(key, default)


def _theme_bool(theme, section, key):
    return _theme_get(theme, section, key, "false") == "true"


def _theme_align(v):
    if v == "center":
        return WD_ALIGN_PARAGRAPH.CENTER
    if v == "right":
        return WD_ALIGN_PARAGRAPH.RIGHT
    if v == "left":
        return WD_ALIGN_PARAGRAPH.LEFT
    return None


def _set_style_font(style, theme, section, prefix=""):
    font_name = _theme_get(theme, section, prefix + "font")
    size = _theme_get(theme, section, prefix + "size")
    bold = _theme_bool(theme, section, prefix + "bold")

    if font_name:
        style.font.name = font_name
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        rFonts.set(qn("w:eastAsia"), font_name)
        rFonts.set(qn("w:ascii"), font_name)
        rFonts.set(qn("w:hAnsi"), font_name)
        for attr in ["w:asciiTheme", "w:eastAsiaTheme", "w:hAnsiTheme", "w:cstheme"]:
            qattr = qn(attr)
            if qattr in rFonts.attrib:
                del rFonts.attrib[qattr]

    if size:
        style.font.size = _theme_size(size)
    if bold:
        style.font.bold = True

    color = _theme_get(theme, section, prefix + "color")
    if not color:
        color = _theme_get(theme, section, "color")
    if color:
        style.font.color.rgb = RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))


def _set_style_paragraph(style, theme, section):
    for key, attr in [("before_spacing", "space_before"), ("after_spacing", "space_after")]:
        v = _theme_get(theme, section, key)
        if v:
            setattr(style.paragraph_format, attr, _theme_size(v))
    v = _theme_get(theme, section, "line_spacing")
    if v:
        style.paragraph_format.line_spacing = float(v)
    v = _theme_get(theme, section, "first_line_indent")
    if v:
        style.paragraph_format.first_line_indent = _theme_size(v)
    v = _theme_get(theme, section, "align")
    if v:
        al = _theme_align(v)
        if al:
            style.paragraph_format.alignment = al


class ThemedDocument:
    """Wrapper that adds content using the style system after a theme is applied."""

    def __init__(self, doc, theme):
        self.doc = doc
        self.theme = theme
        self._fig_num = 0
        self._tbl_num = 0

    def heading(self, text, level=1):
        """Add a heading using the document's style system."""
        return self.doc.add_paragraph(text, style=f"Heading {level}")

    def body(self, text):
        """Add a body paragraph using the Normal style."""
        return self.doc.add_paragraph(text, style="Normal")

    def figure(self, image_path, caption_text):
        """Insert a centered figure with numbered caption."""
        self._fig_num += 1
        prefix = _theme_get(self.theme, "figure", "caption_prefix", "图")

        p_img = self.doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(image_path, width=Cm(10))

        p_cap = self.doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

        r_label = p_cap.add_run(f"{prefix}{self._fig_num} ")
        lbl_font = _theme_get(self.theme, "figure", "label_font")
        lbl_size = _theme_get(self.theme, "figure", "label_size")
        lbl_bold = _theme_bool(self.theme, "figure", "label_bold")
        if lbl_font:
            r_label.font.name = lbl_font
            rPr = r_label._element.get_or_add_rPr()
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rPr.insert(0, rFonts)
            rFonts.set(qn("w:eastAsia"), lbl_font)
        if lbl_size:
            r_label.font.size = _theme_size(lbl_size)
        if lbl_bold:
            r_label.bold = True

        r_text = p_cap.add_run(caption_text)
        cap_font = _theme_get(self.theme, "figure", "caption_font")
        cap_size = _theme_get(self.theme, "figure", "caption_size")
        if cap_font:
            r_text.font.name = cap_font
            rPr = r_text._element.get_or_add_rPr()
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rPr.insert(0, rFonts)
            rFonts.set(qn("w:eastAsia"), cap_font)
        if cap_size:
            r_text.font.size = _theme_size(cap_size)

        p_cap.paragraph_format.space_after = Pt(12)
        return p_cap

    def table(self, headers, rows_data, caption_text):
        """Insert a table with numbered caption above."""
        self._tbl_num += 1
        prefix = _theme_get(self.theme, "table", "caption_prefix", "表")

        p_cap = self.doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_label = p_cap.add_run(f"{prefix}{self._tbl_num} ")
        lbl_font = _theme_get(self.theme, "table", "label_font")
        lbl_size = _theme_get(self.theme, "table", "label_size")
        lbl_bold = _theme_bool(self.theme, "table", "label_bold")
        if lbl_font:
            r_label.font.name = lbl_font
            rPr = r_label._element.get_or_add_rPr()
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rPr.insert(0, rFonts)
            rFonts.set(qn("w:eastAsia"), lbl_font)
        if lbl_size:
            r_label.font.size = _theme_size(lbl_size)
        if lbl_bold:
            r_label.bold = True

        r_text = p_cap.add_run(caption_text)
        cap_font = _theme_get(self.theme, "table", "caption_font")
        cap_size = _theme_get(self.theme, "table", "caption_size")
        if cap_font:
            r_text.font.name = cap_font
            rPr = r_text._element.get_or_add_rPr()
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rPr.insert(0, rFonts)
            rFonts.set(qn("w:eastAsia"), cap_font)
        if cap_size:
            r_text.font.size = _theme_size(cap_size)

        ncols, nrows = len(headers), len(rows_data) + 1
        tbl = self.doc.add_table(rows=nrows, cols=ncols, style="Light Grid Accent 1")

        cfont = _theme_get(self.theme, "table", "content_font")
        csize = _theme_get(self.theme, "table", "content_size")

        for i, h in enumerate(headers):
            cell = tbl.rows[0].cells[i]
            cell.text = h
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.bold = True
                    if cfont:
                        run.font.name = cfont
                    if csize:
                        run.font.size = _theme_size(csize)

        for ri, row_data in enumerate(rows_data):
            for ci, cell_text in enumerate(row_data):
                cell = tbl.rows[ri + 1].cells[ci]
                cell.text = cell_text
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        if cfont:
                            run.font.name = cfont
                        if csize:
                            run.font.size = _theme_size(csize)

        return tbl

    def formula(self, latex_str):
        """Insert a centered LaTeX formula."""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        omml = latex_to_omml(latex_str)
        run = p.add_run("")
        run._element.append(etree.fromstring(omml))
        return p

    def cover(self, title_text, info_lines):
        """Add a cover page with centered title and info lines."""
        for _ in range(6):
            self.doc.add_paragraph("")

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title_text)
        cover_font = _theme_get(self.theme, "cover", "title_font")
        cover_size = _theme_get(self.theme, "cover", "title_size")
        cover_bold = _theme_bool(self.theme, "cover", "title_bold")
        if cover_font:
            r.font.name = cover_font
            rPr = r._element.get_or_add_rPr()
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rPr.insert(0, rFonts)
            rFonts.set(qn("w:eastAsia"), cover_font)
        if cover_size:
            r.font.size = _theme_size(cover_size)
        if cover_bold:
            r.bold = True

        info_font = _theme_get(self.theme, "cover", "info_font")
        info_size = _theme_get(self.theme, "cover", "info_size")

        for line in info_lines:
            self.doc.add_paragraph("")
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line)
            if info_font:
                r.font.name = info_font
            if info_size:
                r.font.size = _theme_size(info_size)


# ═══════════════════════════════════════════════════════════════
# DocxBuilder — CREATE (from docx_tool.py)
# ═══════════════════════════════════════════════════════════════

class DocxBuilder:
    """Fluent API for building .docx documents with Chinese/English typography presets."""

    def __init__(self, template_path=None):
        """Load template or create blank document."""
        self.doc = Document(template_path) if template_path else Document()

    def _remove_theme_refs(self, style):
        """Remove theme font references from a style."""
        rPr = style.element.find(f"{{{_W_NS}}}rPr")
        if rPr is None:
            return
        rFonts = rPr.find(f"{{{_W_NS}}}rFonts")
        if rFonts is None:
            return
        for attr in [
            f"{{{_W_NS}}}eastAsiaTheme",
            f"{{{_W_NS}}}asciiTheme",
            f"{{{_W_NS}}}hAnsiTheme",
            f"{{{_W_NS}}}cstheme",
        ]:
            if attr in rFonts.attrib:
                del rFonts.attrib[attr]

    def _set_style_color(self, style, hex_color="000000"):
        """Set font color on a style (hex without #).
        
        Also removes w:themeColor so Word uses the explicit color instead of theme color.
        Word prioritizes themeColor over val — if themeColor is present, the heading stays blue.
        """
        rPr = style.element.find(f"{{{_W_NS}}}rPr")
        if rPr is None:
            rPr = etree.SubElement(style.element, f"{{{_W_NS}}}rPr")
        color = rPr.find(f"{{{_W_NS}}}color")
        if color is None:
            color = etree.SubElement(rPr, f"{{{_W_NS}}}color")
        color.set(f"{{{_W_NS}}}val", hex_color)
        # CRITICAL: Remove themeColor so Word doesn't override with blue accent color
        for attr in [f"{{{_W_NS}}}themeColor", f"{{{_W_NS}}}themeTint", f"{{{_W_NS}}}themeShade"]:
            if attr in color.attrib:
                del color.attrib[attr]

    def _set_style_east_asian_font(self, style, ea_font, latin_font="Times New Roman"):
        """Set East Asian and Latin fonts on a style."""
        rPr = style.element.find(f"{{{_W_NS}}}rPr")
        if rPr is None:
            rPr = etree.SubElement(style.element, f"{{{_W_NS}}}rPr")
        rFonts = rPr.find(f"{{{_W_NS}}}rFonts")
        if rFonts is None:
            rFonts = etree.SubElement(rPr, f"{{{_W_NS}}}rFonts")
        rFonts.set(f"{{{_W_NS}}}eastAsia", ea_font)
        rFonts.set(f"{{{_W_NS}}}ascii", latin_font)
        rFonts.set(f"{{{_W_NS}}}hAnsi", latin_font)
        style.font.name = latin_font

    def preset_chinese(
        self,
        body_font="SimHei",
        heading_font="SimSun",
        body_size=10.5,
        h1_size=18,
        h2_size=14,
        h3_size=12,
        color="000000",
    ):
        """Apply Chinese document defaults.

        Sets Normal to ``body_font`` at ``body_size`` pt and headings 1–3 to
        ``heading_font`` with bold.  Removes theme font references and sets
        explicit black colour.
        """
        style = self.doc.styles["Normal"]
        style.font.size = Pt(body_size)  # type: ignore[union-attr]
        self._set_style_east_asian_font(style, body_font)
        self._set_style_color(style, color)
        self._remove_theme_refs(style)

        for lvl, sz in [(1, h1_size), (2, h2_size), (3, h3_size)]:
            try:
                s = self.doc.styles[f"Heading {lvl}"]
                s.font.size = Pt(sz)  # type: ignore[union-attr]
                s.font.bold = True  # type: ignore[union-attr]
                self._set_style_east_asian_font(s, heading_font)
                self._set_style_color(s, color)
                self._remove_theme_refs(s)
            except KeyError:
                pass
        return self

    def preset_english(
        self,
        body_font="Times New Roman",
        heading_font="Times New Roman",
        body_size=12,
        h1_size=18,
        h2_size=14,
        h3_size=12,
        color="000000",
    ):
        """Apply English document defaults."""
        style = self.doc.styles["Normal"]
        style.font.size = Pt(body_size)  # type: ignore[union-attr]
        style.font.name = body_font  # type: ignore[union-attr]
        self._set_style_color(style, color)
        self._remove_theme_refs(style)

        for lvl, sz in [(1, h1_size), (2, h2_size), (3, h3_size)]:
            try:
                s = self.doc.styles[f"Heading {lvl}"]
                s.font.size = Pt(sz)  # type: ignore[union-attr]
                s.font.bold = True  # type: ignore[union-attr]
                s.font.name = heading_font  # type: ignore[union-attr]
                self._set_style_color(s, color)
                self._remove_theme_refs(s)
            except KeyError:
                pass
        return self

    def add_heading(self, text, level=1, alignment=None, font=None, color=None):
        """Add a heading paragraph."""
        p = self.doc.add_heading(text, level=level)
        if alignment is not None:
            p.alignment = alignment
        if font or color:
            for run in p.runs:
                if font:
                    run.font.name = font
                if color:
                    run.font.color.rgb = RGBColor(
                        int(color[:2], 16), int(color[2:4], 16), int(color[4:6], 16)
                    )
        return self

    def add_paragraph(self, text, style="Normal", alignment=None, font=None, color=None):
        """Add a body paragraph."""
        p = self.doc.add_paragraph(text, style=style)
        if alignment is not None:
            p.alignment = alignment
        if font or color:
            for run in p.runs:
                if font:
                    run.font.name = font
                if color:
                    run.font.color.rgb = RGBColor(
                        int(color[:2], 16), int(color[2:4], 16), int(color[4:6], 16)
                    )
        return self

    def add_formula(self, latex_str, display=True):
        """Insert a LaTeX formula as native OMML.

        ``display=True``  → block-level (``w:p > m:oMathPara``)
        ``display=False`` → inline (``w:r > m:oMath``)
        """
        p = self.doc.add_paragraph()
        if display:
            omml_xml = latex_to_omml(latex_str, inline=False)
            p._element.append(etree.fromstring(omml_xml))
        else:
            omml_xml = latex_to_omml(latex_str, inline=True)
            run = p.add_run("")
            run._element.append(etree.fromstring(omml_xml))
        return self

    def add_table(self, rows, cols, data=None, style="Table Grid"):
        """Add a table.  ``data`` is a list of lists, each inner list is a row."""
        table = self.doc.add_table(rows=rows, cols=cols, style=style)
        if data:
            for i, row_data in enumerate(data):
                if i >= rows:
                    break
                row = table.rows[i]
                for j, cell_text in enumerate(row_data):
                    if j >= cols:
                        break
                    row.cells[j].text = str(cell_text)
        return self

    def add_page_break(self):
        """Add a page break."""
        self.doc.add_page_break()
        return self

    def add_image(self, image_path, width=None, height=None):
        """Add an image.  width/height in inches (use ``docx.shared.Inches``)."""
        kwargs = {}
        if width:
            kwargs["width"] = width
        if height:
            kwargs["height"] = height
        self.doc.add_picture(image_path, **kwargs)
        return self

    def save(self, path):
        """Save the document."""
        self.doc.save(path)
        return self

    @property
    def paragraph_count(self):
        """Number of paragraphs in the document."""
        return len(self.doc.paragraphs)

    @property
    def table_count(self):
        """Number of tables in the document."""
        return len(self.doc.tables)


# Backward-compatible alias
DocumentBuilder = DocxBuilder


# ═══════════════════════════════════════════════════════════════
# DocxReader — READ (from extract_docx.py)
# ═══════════════════════════════════════════════════════════════

_HEADING_PATTERNS = [
    (r"^第[一二三四五六七八九十\d]+章", 1, 0.95),
    (r"^第[一二三四五六七八九十\d]+节", 2, 0.95),
    (r"^\d+\.\d+\s+\w", 2, 0.9),
    (r"^\d+\.\s+\w", 1, 0.85),
    (r"^\d+\s+\w", 1, 0.8),
    (r"^[一二三四五六七八九十]、", 1, 0.8),
    (
        r"^(Abstract|摘要|Introduction|引言|Conclusion|结论|References|参考文献|Acknowledgements|致谢)$",
        1,
        0.75,
    ),
]


def _get_body_order(doc):
    """Return list of ("para", idx) or ("table", idx) in document order."""
    order = []
    para_idx = 0
    table_idx = 0
    for child in doc.element.body:
        tag = child.tag
        if tag.endswith("}p"):
            order.append(("para", para_idx))
            para_idx += 1
        elif tag.endswith("}tbl"):
            order.append(("table", table_idx))
            table_idx += 1
    return order


def _truncate(text, max_len):
    if len(text) > max_len:
        return text[:max_len] + "..."
    return text


class DocxReader:
    """Read and extract content from an existing .docx file."""

    def __init__(self, path):
        """Open the document at ``path``."""
        self.path = path
        self.doc = Document(path)

    # ── Extraction ─────────────────────────────────────────────

    def extract_to_markdown(self, output_path=None):
        """Extract the entire document to Markdown.

        If ``output_path`` is omitted, writes to ``<input>_extracted.md``.
        Returns the output path.
        """
        if output_path is None:
            base = os.path.splitext(self.path)[0]
            output_path = f"{base}_extracted.md"

        with open(output_path, "w", encoding="utf-8") as f:
            f.write(f"# Source: {os.path.basename(self.path)}\n\n")

            for para in self.doc.paragraphs:
                style = (para.style.name or "") if para.style else "Normal"
                text = para.text.strip()
                if not text:
                    f.write("\n")
                    continue
                if style.startswith("Heading"):
                    try:
                        level = int(style.replace("Heading ", ""))
                        f.write(f"{'#' * level} {text}\n\n")
                    except ValueError:
                        f.write(f"# {text}\n\n")
                else:
                    f.write(f"{text}\n\n")

            if self.doc.tables:
                f.write("\n---\n\n")
            for i, table in enumerate(self.doc.tables):
                f.write(f"\n**Table {i + 1}:**\n\n")
                for row in table.rows:
                    cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
                    f.write("| " + " | ".join(cells) + " |\n")
                f.write("\n")

        print(f"Extracted: {self.path} -> {output_path}")
        print(f"  Paragraphs: {len(self.doc.paragraphs)}")
        print(f"  Tables: {len(self.doc.tables)}")
        return output_path

    def extract_range(self, start, end, output_path=None):
        """Extract paragraphs ``start`` through ``end`` (1-based, inclusive) to Markdown.

        Returns the output path.
        """
        total = len(self.doc.paragraphs)
        if start < 1 or end > total or start > end:
            raise ValueError(f"Invalid range {start}-{end} (document has {total} paragraphs)")

        if output_path is None:
            base = os.path.splitext(self.path)[0]
            output_path = f"{base}_range_{start}_{end}.md"

        with open(output_path, "w", encoding="utf-8") as f:
            f.write(f"# Source: {os.path.basename(self.path)} (paragraphs {start}-{end})\n\n")

            for para in self.doc.paragraphs[start - 1 : end]:
                style = (para.style.name or "") if para.style else "Normal"
                text = para.text.strip()
                if not text:
                    f.write("\n")
                    continue
                if style.startswith("Heading"):
                    try:
                        level = int(style.replace("Heading ", ""))
                        f.write(f"{'#' * level} {text}\n\n")
                    except ValueError:
                        f.write(f"# {text}\n\n")
                else:
                    f.write(f"{text}\n\n")

            body_order = _get_body_order(self.doc)
            first_body_pos = None
            last_body_pos = None
            for pos, (etype, eidx) in enumerate(body_order):
                if etype == "para" and start - 1 <= eidx < end:
                    if first_body_pos is None:
                        first_body_pos = pos
                    last_body_pos = pos

            tables_in_range = []
            if first_body_pos is not None and last_body_pos is not None:
                for pos, (etype, eidx) in enumerate(body_order):
                    if etype == "table" and first_body_pos < pos < last_body_pos:
                        tables_in_range.append(eidx)

            if tables_in_range:
                f.write("\n---\n\n")
            for ti in tables_in_range:
                table = self.doc.tables[ti]
                f.write(f"\n**Table {ti + 1}:**\n\n")
                for row in table.rows:
                    cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
                    f.write("| " + " | ".join(cells) + " |\n")
                f.write("\n")

        print(f"Extracted: {self.path}[{start}-{end}] -> {output_path}")
        return output_path

    def extract_section(self, section_name, output_path=None, exact=False):
        """Extract all paragraphs under heading(s) matching ``section_name`` to Markdown.

        ``exact=True`` requires an exact heading text match.
        Returns the output path.
        """
        structure = self._extract_structure()

        matched = []
        for pi, level, text, _confidence, _reason in structure:
            if exact:
                if text == section_name:
                    matched.append((pi, level, text))
            else:
                if section_name.lower() in text.lower():
                    matched.append((pi, level, text))

        if not matched:
            raise ValueError(f"No section matching '{section_name}'")

        if output_path is None:
            safe_name = re.sub(r"[^\w]", "_", section_name)[:30]
            base = os.path.splitext(self.path)[0]
            output_path = f"{base}_section_{safe_name}.md"

        with open(output_path, "w", encoding="utf-8") as f:
            f.write(f"# Source: {os.path.basename(self.path)} (section: {section_name})\n\n")

            for match_idx, (start_pi, match_level, match_text) in enumerate(matched):
                if match_idx > 0:
                    f.write("\n---\n\n")

                end_pi = len(self.doc.paragraphs) - 1
                for pi, level, _text, _confidence, _reason in structure:
                    if pi > start_pi and level <= match_level:
                        end_pi = pi - 1
                        break

                for pi in range(start_pi, end_pi + 1):
                    para = self.doc.paragraphs[pi]
                    style = (para.style.name or "") if para.style else "Normal"
                    text = para.text.strip()
                    if not text:
                        f.write("\n")
                        continue
                    if style.startswith("Heading"):
                        try:
                            level = int(style.replace("Heading ", ""))
                            f.write(f"{'#' * level} {text}\n\n")
                        except ValueError:
                            f.write(f"# {text}\n\n")
                    else:
                        f.write(f"{text}\n\n")

                body_order = _get_body_order(self.doc)
                para_body_pos = {}
                table_body_pos = {}
                for pos, (etype, eidx) in enumerate(body_order):
                    if etype == "para":
                        para_body_pos[eidx] = pos
                    else:
                        table_body_pos[eidx] = pos

                first_body_pos = None
                last_body_pos = None
                for pos, (etype, eidx) in enumerate(body_order):
                    if etype == "para" and start_pi <= eidx <= end_pi:
                        if first_body_pos is None:
                            first_body_pos = pos
                        last_body_pos = pos

                tables_in_section = []
                if first_body_pos is not None and last_body_pos is not None:
                    for pos, (etype, eidx) in enumerate(body_order):
                        if etype == "table" and first_body_pos < pos < last_body_pos:
                            tables_in_section.append(eidx)

                if tables_in_section:
                    f.write("\n---\n\n")
                for ti in tables_in_section:
                    table = self.doc.tables[ti]
                    f.write(f"\n**Table {ti + 1}:**\n\n")
                    for row in table.rows:
                        cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
                        f.write("| " + " | ".join(cells) + " |\n")
                    f.write("\n")

        print(f"Extracted: {self.path}[section='{section_name}'] -> {output_path}")
        print(f"  Matched {len(matched)} heading(s)")
        return output_path

    # ── Structure / Search ─────────────────────────────────────

    def _extract_structure(self, detect_freeform=False):
        """Internal: extract heading structure."""
        structures = []
        style_detected = set()

        for i, para in enumerate(self.doc.paragraphs):
            style = (para.style.name or "") if para.style else ""
            if style.startswith("Heading"):
                try:
                    level = int(style.replace("Heading ", ""))
                except ValueError:
                    level = 1
                text = para.text.strip()
                if text:
                    structures.append((i, level, text, 1.0, f"style:{style}"))
                    style_detected.add(i)

        if detect_freeform:
            for i, para in enumerate(self.doc.paragraphs):
                if i in style_detected:
                    continue
                text = para.text.strip()
                if not text:
                    continue
                for pattern, level, confidence in _HEADING_PATTERNS:
                    if re.match(pattern, text):
                        if confidence < 0.9:
                            if i + 1 >= len(self.doc.paragraphs):
                                continue
                            next_text = self.doc.paragraphs[i + 1].text.strip()
                            if len(next_text) <= len(text):
                                continue
                            if len(text.split()) > 10:
                                continue
                        structures.append((i, level, text, confidence, f"pattern:{pattern}"))
                        break

        structures.sort(key=lambda x: x[0])
        return structures

    def print_structure(self, detect_freeform=False):
        """Print the document's heading tree to stdout."""
        structure = self._extract_structure(detect_freeform=detect_freeform)
        para_count = len(self.doc.paragraphs)
        table_count = len(self.doc.tables)
        print(f"[document] {os.path.basename(self.path)} ({para_count} paragraphs, {table_count} tables)")

        if not structure:
            return

        body_order = _get_body_order(self.doc)
        para_body_pos = {}
        table_body_pos = {}
        for pos, (etype, eidx) in enumerate(body_order):
            if etype == "para":
                para_body_pos[eidx] = pos
            else:
                table_body_pos[eidx] = pos

        n = len(structure)
        for idx, (pi, level, text, confidence, reason) in enumerate(structure):
            if idx + 1 < n:
                end_pi = structure[idx + 1][0] - 1
            else:
                end_pi = len(self.doc.paragraphs) - 1

            heading_body_pos = para_body_pos.get(pi, -1)
            if idx + 1 < n:
                next_pi = structure[idx + 1][0]
                next_body_pos = para_body_pos.get(next_pi, len(body_order))
            else:
                next_body_pos = len(body_order)

            tables_in_range = []
            for ti, tpos in table_body_pos.items():
                if heading_body_pos < tpos < next_body_pos:
                    tables_in_range.append(ti)

            range_parts = [f"p{pi}-p{end_pi}"]
            for ti in tables_in_range:
                range_parts.append(f"table:{ti}")
            range_str = ", ".join(range_parts)

            indent = "|   " * (level - 1)
            is_freeform = confidence < 1.0
            if is_freeform:
                heading_label = f"H{level}"
                conf_str = f" [conf:{confidence}]"
            else:
                heading_label = f"Heading {level}"
                conf_str = ""

            print(f"{indent}+-- {heading_label} (p{pi}){conf_str}: {text} [{range_str}]")

    def grep(self, pattern, context=0):
        """Search paragraphs for ``pattern`` and print matches with optional context."""
        paragraphs = self.doc.paragraphs
        matches = []
        for i, para in enumerate(paragraphs):
            if re.search(pattern, para.text, re.IGNORECASE):
                matches.append(i)

        if not matches:
            print(f"No matches found for '{pattern}'")
            return

        for idx, pi in enumerate(matches):
            if idx > 0:
                print()
            text = paragraphs[pi].text.strip()
            display = _truncate(text, 120)
            print(f'[paragraph {pi}] "{display}"')

            for offset in range(1, context + 1):
                prev_pi = pi - offset
                if prev_pi >= 0:
                    prev_text = paragraphs[prev_pi].text.strip()
                    prev_display = _truncate(prev_text, 120)
                    print(f'  <- p{prev_pi}: "{prev_display}"')

            for offset in range(1, context + 1):
                next_pi = pi + offset
                if next_pi < len(paragraphs):
                    next_text = paragraphs[next_pi].text.strip()
                    next_display = _truncate(next_text, 120)
                    print(f'  -> p{next_pi}: "{next_display}"')


# ═══════════════════════════════════════════════════════════════
# DocxEditor — EDIT (from edit_docx.py)
# ═══════════════════════════════════════════════════════════════

def _set_paragraph_text(para, new_text):
    """Replace entire paragraph text, preserving formatting of the first run."""
    if not para.runs:
        para.add_run(new_text)
        return
    for run in para.runs[1:]:
        run._element.getparent().remove(run._element)
    para.runs[0].text = new_text


def _replace_in_paragraph(para, old, new):
    """Replace text in a paragraph, handling text split across multiple runs."""
    full_text = para.text
    if old not in full_text:
        return False
    for run in para.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return True
    _set_paragraph_text(para, full_text.replace(old, new))
    return True


def _replace_in_table_cell(cell, old, new):
    changed = False
    for para in cell.paragraphs:
        if _replace_in_paragraph(para, old, new):
            changed = True
    return changed


def _replace_all(doc, old, new):
    """Replace text across entire document (paragraphs, tables, headers, footers)."""
    count = 0
    for para in doc.paragraphs:
        if _replace_in_paragraph(para, old, new):
            count += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if _replace_in_table_cell(cell, old, new):
                    count += 1
    for section in doc.sections:
        for hf in [
            section.header,
            section.footer,
            section.even_page_header,
            section.even_page_footer,
            section.first_page_header,
            section.first_page_footer,
        ]:
            if hf is None:
                continue
            for para in hf.paragraphs:
                if _replace_in_paragraph(para, old, new):
                    count += 1
            for table in hf.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if _replace_in_table_cell(cell, old, new):
                            count += 1
    return count


class DocxEditor:
    """Edit an existing .docx document with robust text replacement and structural edits."""

    def __init__(self, path):
        """Open the document at ``path`` for editing."""
        self.path = path
        self.doc = Document(path)

    # ── Text replacement ───────────────────────────────────────

    def replace(self, old, new):
        """Replace ``old`` with ``new`` across the entire document.

        Returns the number of paragraphs/cells modified.
        """
        count = _replace_all(self.doc, old, new)
        print(f"Replace '{old}' -> '{new}': {count} occurrences")
        return count

    def replace_in_range(self, start, end, old, new):
        """Replace ``old`` with ``new`` only within paragraph range ``start``–``end`` (inclusive, 0-based).

        Returns the number of paragraphs modified.
        """
        count = 0
        for idx in range(start, end + 1):
            if _replace_in_paragraph(self.doc.paragraphs[idx], old, new):
                count += 1
        print(f"Range [{start}-{end}]: replaced '{old}' -> '{new}': {count} occurrences")
        return count

    # ── Structural edits ───────────────────────────────────────

    def insert_after(self, index, text, style=None):
        """Insert a new paragraph after paragraph ``index`` (0-based)."""
        if index < 0 or index >= len(self.doc.paragraphs):
            raise IndexError(f"Paragraph index {index} out of range (0-{len(self.doc.paragraphs) - 1})")
        target = self.doc.paragraphs[index]
        new_para = self.doc.add_paragraph(text)
        if style:
            new_para.style = style
        target._element.addnext(new_para._element)
        print(f"Inserted paragraph after index {index}")
        return new_para

    def delete(self, index):
        """Delete paragraph at ``index`` (0-based)."""
        if index < 0 or index >= len(self.doc.paragraphs):
            raise IndexError(f"Paragraph index {index} out of range (0-{len(self.doc.paragraphs) - 1})")
        para = self.doc.paragraphs[index]
        para._element.getparent().remove(para._element)
        print(f"Deleted paragraph {index}: '{para.text[:50]}...'")

    def set_style(self, index, style_name):
        """Set paragraph ``index`` (0-based) to ``style_name``."""
        if index < 0 or index >= len(self.doc.paragraphs):
            raise IndexError(f"Paragraph index {index} out of range (0-{len(self.doc.paragraphs) - 1})")
        self.doc.paragraphs[index].style = self.doc.styles[style_name]
        print(f"Set paragraph {index} style to '{style_name}'")

    def set_paragraph_text(self, index, text):
        """Replace the entire text of paragraph ``index`` (0-based)."""
        if index < 0 or index >= len(self.doc.paragraphs):
            raise IndexError(f"Paragraph index {index} out of range (0-{len(self.doc.paragraphs) - 1})")
        _set_paragraph_text(self.doc.paragraphs[index], text)
        print(f"Set paragraph {index} to '{text}'")

    def append_to_paragraph(self, index, text):
        """Append ``text`` to the end of paragraph ``index``."""
        para = self.doc.paragraphs[index]
        _set_paragraph_text(para, para.text + " " + text)
        print(f"Appended to paragraph {index}")

    def prepend_to_paragraph(self, index, text):
        """Prepend ``text`` to the start of paragraph ``index``."""
        para = self.doc.paragraphs[index]
        _set_paragraph_text(para, text + " " + para.text)
        print(f"Prepended to paragraph {index}")

    def set_paragraph_from_file(self, index, file_path):
        """Replace paragraph ``index`` with the contents of ``file_path``."""
        with open(file_path, "r", encoding="utf-8") as f:
            new_text = f.read()
        _set_paragraph_text(self.doc.paragraphs[index], new_text)
        print(f"Set paragraph {index} to content of '{file_path}'")

    # ─-- Track Changes ─────────────────────────────────────────

    def add_tracked_insertion(self, index, text, author="Editor"):
        """Add tracked insertion text at the end of paragraph ``index``."""
        from datetime import datetime

        para = self.doc.paragraphs[index]
        date_str = datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ")
        ins = etree.SubElement(para._element, qn("w:ins"))
        ins.set(qn("w:id"), str(os.urandom(2)[0] % 1000))
        ins.set(qn("w:author"), author)
        ins.set(qn("w:date"), date_str)
        r = etree.SubElement(ins, qn("w:r"))
        rPr = etree.SubElement(r, qn("w:rPr"))
        t = etree.SubElement(r, qn("w:t"))
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = text
        print(f"Added tracked insertion to paragraph {index}")

    def add_tracked_deletion(self, index, author="Editor"):
        """Mark paragraph ``index`` as a tracked deletion."""
        from datetime import datetime

        para = self.doc.paragraphs[index]
        date_str = datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ")
        pPr = para._element.find(qn("w:pPr"))
        if pPr is None:
            pPr = etree.SubElement(para._element, qn("w:pPr"))
            para._element.insert(0, pPr)
        rPr = pPr.find(qn("w:rPr"))
        if rPr is None:
            rPr = etree.SubElement(pPr, qn("w:rPr"))
        del_mark = etree.SubElement(rPr, qn("w:del"))
        del_mark.set(qn("w:id"), str(os.urandom(2)[0] % 1000))
        del_mark.set(qn("w:author"), author)
        del_mark.set(qn("w:date"), date_str)

        for run_elem in list(para._element.findall(qn("w:r"))):
            del_elem = etree.Element(qn("w:del"))
            del_elem.set(qn("w:id"), str(os.urandom(2)[0] % 1000))
            del_elem.set(qn("w:author"), author)
            del_elem.set(qn("w:date"), date_str)
            para._element.replace(run_elem, del_elem)
            for t_elem in run_elem.findall(qn("w:t")):
                t_elem.tag = qn("w:delText")
            del_elem.append(run_elem)
        print(f"Marked paragraph {index} as tracked deletion")

    # ── Fill from JSON ─────────────────────────────────────────

    def fill_from_json(self, json_path):
        """Apply replacements from a JSON file mapping old→new strings."""
        with open(json_path, "r", encoding="utf-8") as f:
            replacements = json.load(f)
        for old, new in replacements.items():
            count = _replace_all(self.doc, old, new)
            print(f"Replace '{old}' -> '{new}': {count} occurrences")

    # ── List paragraphs ────────────────────────────────────────

    def list_paragraphs(self, max_show=None):
        """Print all paragraphs with indices (useful for finding edit targets)."""
        total = len(self.doc.paragraphs)
        for i, para in enumerate(self.doc.paragraphs):
            text = para.text.strip()
            if max_show and i >= max_show:
                break
            style = para.style.name if para.style else "-"
            preview = text[:80] + ("..." if len(text) > 80 else "")
            print(f"[{i:3d}] [{style:20s}] {preview}")
        if max_show and total > max_show:
            print(f"... and {total - max_show} more paragraphs")

    # ── Save ───────────────────────────────────────────────────

    def save(self, path=None):
        """Save the edited document.  If ``path`` is omitted, overwrites the original."""
        out = path or self.path
        self.doc.save(out)
        print(f"Saved: {out}")
        return out


# ═══════════════════════════════════════════════════════════════
# DocxUtil — UTILITIES (merge_runs, apply_theme, convert_to_doc)
# ═══════════════════════════════════════════════════════════════

class DocxUtil:
    """Static utility methods for .docx manipulation."""

    # ── merge_runs ─────────────────────────────────────────────

    @staticmethod
    def merge_runs(input_path, output_path):
        """Unpack a .docx, merge adjacent runs with identical formatting, and repack.

        This eliminates the #1 editing bug: text split across multiple ``<w:r>`` elements.
        """
        tmpdir = tempfile.mkdtemp()
        try:
            unpacked = os.path.join(tmpdir, "unpacked")
            os.makedirs(unpacked)
            with zipfile.ZipFile(input_path, "r") as z:
                z.extractall(unpacked)

            count = DocxUtil._merge_runs_in_dir(unpacked)

            with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as z:
                for root, _dirs, files in os.walk(unpacked):
                    for fn in files:
                        fp = os.path.join(root, fn)
                        arcname = os.path.relpath(fp, unpacked)
                        z.write(fp, arcname)

            print(f"Merged {count} runs: {input_path} -> {output_path}")
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)

    @staticmethod
    def _merge_runs_in_dir(unpacked_dir):
        ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        doc_xml = Path(unpacked_dir) / "word" / "document.xml"
        if not doc_xml.exists():
            return 0

        tree = etree.parse(str(doc_xml))
        root = tree.getroot()
        nsmap = {"w": ns}

        for elem in root.findall(".//w:proofErr", nsmap):
            elem.getparent().remove(elem)

        for run in root.findall(".//w:r", nsmap):
            for attr_name in list(run.attrib.keys()):
                if "rsid" in attr_name.lower():
                    del run.attrib[attr_name]

        containers = set()
        for run in root.findall(".//w:r", nsmap):
            containers.add(run.getparent())

        total_merged = 0
        for container in containers:
            total_merged += DocxUtil._merge_in_container(container, nsmap)

        tree.write(str(doc_xml), xml_declaration=True, encoding="UTF-8", pretty_print=True)
        return total_merged

    @staticmethod
    def _merge_in_container(container, nsmap):
        def _is_run(elem):
            return hasattr(elem, "tag") and etree.QName(elem).localname == "r"

        def _is_tag(elem, tag):
            return hasattr(elem, "tag") and etree.QName(elem).localname == tag

        def _is_text_node(elem):
            return not hasattr(elem, "tag")

        def _is_element(elem):
            return hasattr(elem, "tag")

        def _same_format(run1, run2):
            def _find_child(parent, tag):
                for child in parent:
                    if hasattr(child, "tag") and etree.QName(child).localname == tag:
                        return child
                return None

            rpr1 = _find_child(run1, "rPr")
            rpr2 = _find_child(run2, "rPr")
            if (rpr1 is None) != (rpr2 is None):
                return False
            if rpr1 is None:
                return True
            return etree.tostring(rpr1) == etree.tostring(rpr2)

        def _merge_runs(target, source):
            for child in list(source):
                tag = etree.QName(child).localname if hasattr(child, "tag") else ""
                if tag != "rPr":
                    target.append(child)

        def _consolidate_text(run):
            t_elements = [c for c in run if _is_tag(c, "t")]
            for k in range(len(t_elements) - 1, 0, -1):
                curr = t_elements[k]
                prev = t_elements[k - 1]
                if DocxUtil._adjacent(prev, curr, run):
                    prev_text = prev.text or ""
                    curr_text = curr.text or ""
                    prev.text = prev_text + curr_text
                    if prev.text.startswith(" ") or prev.text.endswith(" "):
                        prev.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    run.remove(curr)

        merged = 0
        children = list(container)
        i = 0
        while i < len(children):
            current = children[i]
            if not _is_run(current):
                i += 1
                continue
            j = i + 1
            while j < len(children):
                next_elem = children[j]
                if _is_run(next_elem) and _same_format(current, next_elem):
                    _merge_runs(current, next_elem)
                    container.remove(next_elem)
                    children = list(container)
                    merged += 1
                elif _is_run(next_elem):
                    break
                elif _is_text_node(next_elem) and (next_elem.tail or "").strip():
                    break
                else:
                    j += 1
            _consolidate_text(current)
            i += 1
        return merged

    @staticmethod
    def _adjacent(elem1, elem2, parent):
        between = False
        for child in parent:
            if child is elem1:
                between = True
            elif child is elem2:
                return between and True
            elif between:
                if not hasattr(child, "tag"):
                    continue
                return False
            if between and hasattr(child, "tag"):
                return False
        return False

    # ── apply_theme ────────────────────────────────────────────

    @staticmethod
    def load_theme(path, degree="硕士"):
        """Parse a format-spec theme file into a nested dict.

        ``degree`` replaces ``{degree}`` placeholders in theme values.
        """
        cfg, section = {}, None
        with open(path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line or line.startswith("#"):
                    continue
                m = re.match(r"^\[(\w+)\]$", line)
                if m:
                    section = m.group(1)
                    cfg[section] = {}
                    continue
                m2 = re.match(r"^(\w+)\s*=\s*(.+)$", line)
                if m2 and section:
                    cfg[section][m2.group(1)] = m2.group(2).strip().replace("{degree}", degree)
        return cfg

    @staticmethod
    def apply_theme(input_path, theme_path, output_path, degree="硕士"):
        """Apply a format-spec theme to a .docx file and save to ``output_path``."""
        theme = DocxUtil.load_theme(theme_path, degree=degree)
        doc = Document(input_path)

        # Page setup
        for sec in doc.sections:
            for key, attr in [
                ("page_width", "page_width"),
                ("page_height", "page_height"),
                ("margin_top", "top_margin"),
                ("margin_bottom", "bottom_margin"),
                ("margin_left", "left_margin"),
                ("margin_right", "right_margin"),
            ]:
                v = _theme_get(theme, "document", key)
                if v:
                    setattr(sec, attr, _theme_size(v))

        # Normal
        normal = doc.styles["Normal"]
        _set_style_font(normal, theme, "body")
        _set_style_paragraph(normal, theme, "body")

        # Headings
        heading_map = {
            "Heading 1": "heading1",
            "Heading 2": "heading2",
            "Heading 3": "heading3",
        }
        for style_name, section in heading_map.items():
            if section in theme:
                style = doc.styles[style_name]
                _set_style_font(style, theme, section)
                _set_style_paragraph(style, theme, section)

        # Header
        sec = doc.sections[0]
        hdr_text = _theme_get(theme, "header", "text")
        if hdr_text:
            header = sec.header
            header.paragraphs[0].text = ""
            run = header.paragraphs[0].add_run(hdr_text)
            font = _theme_get(theme, "header", "font")
            size = _theme_get(theme, "header", "size")
            if font:
                run.font.name = font
                rPr = run._element.get_or_add_rPr()
                rFonts = rPr.find(qn("w:rFonts"))
                if rFonts is None:
                    rFonts = OxmlElement("w:rFonts")
                    rPr.insert(0, rFonts)
                rFonts.set(qn("w:eastAsia"), font)
            if size:
                run.font.size = _theme_size(size)

        # Footer
        footer = sec.footer
        footer.paragraphs[0].text = ""
        footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        ffont = _theme_get(theme, "footer", "font")
        fsize = _theme_get(theme, "footer", "size")
        if ffont or fsize:
            run = footer.paragraphs[0].add_run(" ")
            if ffont:
                run.font.name = ffont
            if fsize:
                run.font.size = _theme_size(fsize)

        doc.save(output_path)
        print(f"Applied theme: {input_path} + {theme_path} -> {output_path}")

    # ── convert_to_doc ─────────────────────────────────────────

    @staticmethod
    def convert_to_doc(docx_path, doc_path=None):
        """Convert ``.docx`` to legacy ``.doc`` format.

        On Windows tries MS Word COM first (requires pywin32), then falls back
        to LibreOffice.  On other platforms uses LibreOffice exclusively.
        """
        if doc_path is None:
            doc_path = os.path.splitext(docx_path)[0] + ".doc"

        if platform.system() == "Windows":
            try:
                DocxUtil._convert_win32(docx_path, doc_path)
                return doc_path
            except ImportError:
                print("pywin32 not installed, trying LibreOffice...")
            except Exception as e:
                print(f"MS Word COM failed ({e}), trying LibreOffice...")

        try:
            DocxUtil._convert_libreoffice(docx_path, doc_path)
        except FileNotFoundError:
            print("Error: LibreOffice (soffice) not found in PATH.")
            print("Install LibreOffice: https://www.libreoffice.org/download/")
            print("Or on Windows with MS Word: pip install pywin32")
            sys.exit(1)
        return doc_path

    @staticmethod
    def _convert_win32(docx_path, doc_path):
        import win32com.client

        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(os.path.abspath(docx_path))
        doc.SaveAs2(os.path.abspath(doc_path), FileFormat=0)
        doc.Close()
        word.Quit()
        print(f"Converted (MS Word): {docx_path} -> {doc_path}")

    @staticmethod
    def _convert_libreoffice(docx_path, doc_path):
        out_dir = os.path.dirname(doc_path) or "."
        cmd = [
            "soffice",
            "--headless",
            "--convert-to",
            "doc",
            "--outdir",
            out_dir,
            docx_path,
        ]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
        if result.returncode != 0:
            print("Error: LibreOffice conversion failed")
            print(result.stderr)
            sys.exit(1)
        libre_out = os.path.join(
            out_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".doc"
        )
        if libre_out != doc_path and os.path.exists(libre_out):
            os.replace(libre_out, doc_path)
        print(f"Converted (LibreOffice): {docx_path} -> {doc_path}")


# ═══════════════════════════════════════════════════════════════
# CLI Entry Point
# ═══════════════════════════════════════════════════════════════


def _cli():
    parser = argparse.ArgumentParser(description="Unified .docx tool — create, read, edit, convert.")
    sub = parser.add_subparsers(dest="cmd", required=True)

    # ── create ─────────────────────────────────────────────────
    p_create = sub.add_parser("create", help="Create a new document")
    p_create.add_argument("--output", "-o", required=True, help="Output .docx path")
    p_create.add_argument(
        "--preset", choices=["chinese", "english"], default="chinese",
        help="Typography preset (default: chinese)",
    )
    p_create.add_argument("--title", help="Document title (Heading 1)")
    p_create.add_argument(
        "--add-heading", action="append", default=[],
        help='Add heading (can repeat). Format: "LEVEL:TEXT" e.g. "1:Overview"',
    )
    p_create.add_argument(
        "--add-paragraph", action="append", default=[], help="Add paragraph (can repeat)",
    )
    p_create.add_argument(
        "--add-formula", action="append", default=[],
        help="Add LaTeX formula as display block (can repeat)",
    )
    p_create.add_argument(
        "--add-inline-formula", action="append", default=[],
        help="Add inline LaTeX formula (can repeat)",
    )
    p_create.add_argument("--body-font", default=None, help="Override body font")
    p_create.add_argument("--heading-font", default=None, help="Override heading font")
    p_create.add_argument("--color", default="000000", help="Font color hex (default: 000000)")

    # ── extract (read) ─────────────────────────────────────────
    p_extract = sub.add_parser("extract", help="Extract/read a .docx")
    p_extract.add_argument("input", help="Input .docx path")
    p_extract.add_argument("--output", "-o", help="Output Markdown path")
    p_extract.add_argument("--structure", action="store_true", help="Print heading tree")
    p_extract.add_argument("--detect-freeform", action="store_true", help="Detect free-form headings")
    p_extract.add_argument("--range", help="Extract paragraph range N-M (1-based)")
    p_extract.add_argument("--section", help="Extract by heading name")
    p_extract.add_argument("--exact", action="store_true", help="Exact heading match for --section")
    p_extract.add_argument("--grep", help="Search paragraphs for pattern")
    p_extract.add_argument("--context", type=int, default=0, help="Context lines for --grep")

    # ── edit ───────────────────────────────────────────────────
    p_edit = sub.add_parser("edit", help="Edit an existing .docx")
    p_edit.add_argument("input", help="Input .docx path")
    p_edit.add_argument("output", help="Output .docx path")
    p_edit.add_argument("--replace", nargs=2, action="append", default=[],
                        metavar=("OLD", "NEW"), help="Replace text (can repeat)")
    p_edit.add_argument("--range-replace", nargs=3, action="append", default=[],
                        metavar=("N-M", "OLD", "NEW"),
                        help="Replace within paragraph range N-M (can repeat)")
    p_edit.add_argument("--insert-after", nargs=2, action="append", default=[],
                        metavar=("INDEX", "TEXT"), help="Insert paragraph after index (can repeat)")
    p_edit.add_argument("--delete", type=int, action="append", default=[],
                        help="Delete paragraph index (can repeat)")
    p_edit.add_argument("--set-style", nargs=2, action="append", default=[],
                        metavar=("INDEX", "STYLE"), help="Set paragraph style (can repeat)")
    p_edit.add_argument("--paragraph", nargs=2, action="append", default=[],
                        metavar=("INDEX", "TEXT"), help="Replace paragraph text (can repeat)")
    p_edit.add_argument("--fill", help="JSON file with replacement mapping")

    # ── merge-runs ─────────────────────────────────────────────
    p_merge = sub.add_parser("merge-runs", help="Merge adjacent runs with identical formatting")
    p_merge.add_argument("input", help="Input .docx path")
    p_merge.add_argument("output", help="Output .docx path")

    # ── convert-to-doc ─────────────────────────────────────────
    p_convert = sub.add_parser("convert-to-doc", help="Convert .docx to .doc")
    p_convert.add_argument("input", help="Input .docx path")
    p_convert.add_argument("output", nargs="?", help="Output .doc path (optional)")

    # ── apply-theme ────────────────────────────────────────────
    p_theme = sub.add_parser("apply-theme", help="Apply a format-spec theme to a .docx")
    p_theme.add_argument("input", help="Input .docx path")
    p_theme.add_argument("theme", help="Theme file path")
    p_theme.add_argument("output", help="Output .docx path")

    args = parser.parse_args()

    if args.cmd == "create":
        builder = DocxBuilder()
        if args.preset == "chinese":
            kw = {}
            if args.body_font:
                kw["body_font"] = args.body_font
            if args.heading_font:
                kw["heading_font"] = args.heading_font
            if args.color:
                kw["color"] = args.color
            builder.preset_chinese(**kw)
        else:
            builder.preset_english()

        if args.title:
            builder.add_heading(args.title, level=1)

        for h in args.add_heading:
            if ":" in h:
                lvl, text = h.split(":", 1)
                builder.add_heading(text, level=int(lvl))
            else:
                builder.add_heading(h, level=1)

        for p in args.add_paragraph:
            builder.add_paragraph(p)

        for f in args.add_formula:
            builder.add_formula(f, display=True)

        for f in args.add_inline_formula:
            builder.add_formula(f, display=False)

        builder.save(args.output)
        print(f"Created: {args.output}")

    elif args.cmd == "extract":
        reader = DocxReader(args.input)

        if args.structure:
            reader.print_structure(detect_freeform=args.detect_freeform)
        elif args.grep:
            reader.grep(args.grep, context=args.context)
        elif args.range:
            m = re.match(r"^(\d+)-(\d+)$", args.range)
            if not m:
                parser.error("Invalid --range format. Use N-M (e.g., 2-5)")
            start, end = int(m.group(1)), int(m.group(2))
            reader.extract_range(start, end, args.output)
        elif args.section:
            reader.extract_section(args.section, args.output, exact=args.exact)
        else:
            reader.extract_to_markdown(args.output)

    elif args.cmd == "edit":
        editor = DocxEditor(args.input)

        if args.fill:
            editor.fill_from_json(args.fill)

        for old, new in args.replace:
            editor.replace(old, new)

        for range_str, old, new in args.range_replace:
            m = re.match(r"^(\d+)-(\d+)$", range_str)
            if not m:
                parser.error(f"Invalid range format '{range_str}'. Expected N-M")
            start, end = int(m.group(1)), int(m.group(2))
            editor.replace_in_range(start, end, old, new)

        for idx_str, text in args.insert_after:
            editor.insert_after(int(idx_str), text)

        for idx in args.delete:
            editor.delete(idx)

        for idx_str, style_name in args.set_style:
            editor.set_style(int(idx_str), style_name)

        for idx_str, text in args.paragraph:
            editor.set_paragraph_text(int(idx_str), text)

        editor.save(args.output)

    elif args.cmd == "merge-runs":
        DocxUtil.merge_runs(args.input, args.output)

    elif args.cmd == "convert-to-doc":
        DocxUtil.convert_to_doc(args.input, args.output)

    elif args.cmd == "apply-theme":
        DocxUtil.apply_theme(args.input, args.theme, args.output)


if __name__ == "__main__":
    _cli()
