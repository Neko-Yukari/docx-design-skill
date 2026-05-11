#!/usr/bin/env python3
"""
Edit .docx files with robust text replacement and structural edits.
Usage: python edit_docx.py <input.docx> <output.docx> [--replace OLD NEW]... [--range N-M --replace OLD NEW]... [--insert-after N TEXT] [--delete N] [--set-style N "StyleName"]... [--paragraph N "TEXT"|--file PATH|--append TEXT|--prepend TEXT]

Examples:
  # Simple find-and-replace
  python edit_docx.py in.docx out.docx --replace "old" "new"

  # Multiple replacements
  python edit_docx.py in.docx out.docx --replace "张三" "李四" --replace "2024" "2025"

  # Replace only within paragraph range 2-5
  python edit_docx.py in.docx out.docx --range 2-5 --replace "天线" "antenna"

  # Insert text after paragraph 3
  python edit_docx.py in.docx out.docx --insert-after 3 "This is a new paragraph."

  # Delete paragraph 5
  python edit_docx.py in.docx out.docx --delete 5

  # Set paragraph style (e.g., "Heading 1", "Normal")
  python edit_docx.py in.docx out.docx --set-style 2 "Heading 1"

  # Multiple style changes
  python edit_docx.py in.docx out.docx --set-style 0 "Heading 1" --set-style 2 "Heading 2"

  # Set paragraph text (replace entire paragraph N)
  python edit_docx.py in.docx out.docx --paragraph 2 "New paragraph text"

  # Replace paragraph text from file content
  python edit_docx.py in.docx out.docx --paragraph 2 --file content.md

  # Append/prepend text to paragraph
  python edit_docx.py in.docx out.docx --paragraph 3 --append " (updated)"
  python edit_docx.py in.docx out.docx --paragraph 3 --prepend "Note: "

  # Combined: fill template blanks
  python edit_docx.py template.docx report.docx --replace "姓名：" "姓名：张三" --replace "学号：" "学号：2024001"
"""
import sys
import os
import re
from docx import Document
from docx.oxml.ns import qn
from copy import deepcopy


# ---------------------------------------------------------------------------
# Core: robust text replace (handles split runs)
# ---------------------------------------------------------------------------

def set_paragraph_text(para, new_text):
    """Replace entire paragraph text, preserving formatting of the first run."""
    if not para.runs:
        para.add_run(new_text)
        return
    # Preserve first run's formatting
    ref_run = para.runs[0]
    # Clear all existing runs
    for run in para.runs[1:]:
        run._element.getparent().remove(run._element)
    ref_run.text = new_text


def replace_in_paragraph(para, old, new):
    """Replace text in a paragraph, handling text split across multiple runs."""
    # Reconstruct full paragraph text
    full_text = para.text
    if old not in full_text:
        return False

    # If simple case (text in a single run), do fast path
    for run in para.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return True

    # Complex case: text spans multiple runs
    # Strategy: set all text into first run, clear others
    set_paragraph_text(para, full_text.replace(old, new))
    return True


def replace_in_table_cell(cell, old, new):
    """Replace text in a table cell."""
    changed = False
    for para in cell.paragraphs:
        if replace_in_paragraph(para, old, new):
            changed = True
    return changed


def replace_all(doc, old, new):
    """Replace text across entire document: paragraphs, tables, headers, footers."""
    count = 0

    # Body paragraphs
    for para in doc.paragraphs:
        if replace_in_paragraph(para, old, new):
            count += 1

    # Tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if replace_in_table_cell(cell, old, new):
                    count += 1

    # Headers & Footers
    for section in doc.sections:
        for hf in [section.header, section.footer,
                   section.even_page_header, section.even_page_footer,
                   section.first_page_header, section.first_page_footer]:
            if hf is None:
                continue
            for para in hf.paragraphs:
                if replace_in_paragraph(para, old, new):
                    count += 1
            for table in hf.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if replace_in_table_cell(cell, old, new):
                            count += 1

    return count


# ---------------------------------------------------------------------------
# Structural edits
# ---------------------------------------------------------------------------

def insert_paragraph_after(doc, after_index, text, style=None):
    """Insert a paragraph after the given paragraph index (0-based)."""
    if after_index < 0 or after_index >= len(doc.paragraphs):
        print(f"Error: paragraph index {after_index} out of range (0-{len(doc.paragraphs)-1})")
        return

    target = doc.paragraphs[after_index]
    new_para = doc.add_paragraph(text)
    if style:
        new_para.style = style
    target._element.addnext(new_para._element)
    print(f"Inserted paragraph after index {after_index}")


def delete_paragraph(doc, index):
    """Delete paragraph at given index (0-based)."""
    if index < 0 or index >= len(doc.paragraphs):
        print(f"Error: paragraph index {index} out of range (0-{len(doc.paragraphs)-1})")
        return

    para = doc.paragraphs[index]
    para._element.getparent().remove(para._element)
    print(f"Deleted paragraph {index}: '{para.text[:50]}...'")


def insert_text_in_paragraph(para, position, text):
    """Insert text at a specific character position within a paragraph."""
    full_text = para.text
    new_text = full_text[:position] + text + full_text[position:]
    set_paragraph_text(para, new_text)


# ---------------------------------------------------------------------------
# Paragraph listing (for finding indices)
# ---------------------------------------------------------------------------

def list_paragraphs(doc, max_show=None):
    """Print all paragraphs with indices, useful for finding positions to edit."""
    total = len(doc.paragraphs)
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if max_show and i >= max_show:
            break
        style = para.style.name if para.style else '-'
        preview = text[:80] + ('...' if len(text) > 80 else '')
        print(f"[{i:3d}] [{style:20s}] {preview}")
    if max_show and total > max_show:
        print(f"... and {total - max_show} more paragraphs")


# ---------------------------------------------------------------------------
# Track Changes (OOXML level)
# ---------------------------------------------------------------------------

def add_tracked_insertion(para, text, author="Editor"):
    """Add text as a tracked insertion at the end of a paragraph."""
    from datetime import datetime
    from lxml import etree

    date_str = datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ")

    # Create <w:ins> element
    nsmap = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    ins = etree.SubElement(para._element, qn('w:ins'))
    ins.set(qn('w:id'), str(os.urandom(2)[0] % 1000))
    ins.set(qn('w:author'), author)
    ins.set(qn('w:date'), date_str)

    # Add a run inside
    r = etree.SubElement(ins, qn('w:r'))
    rPr = etree.SubElement(r, qn('w:rPr'))
    t = etree.SubElement(r, qn('w:t'))
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    t.text = text


def add_tracked_deletion(para, author="Editor"):
    """Mark entire paragraph as tracked deletion."""
    from datetime import datetime
    from lxml import etree

    date_str = datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ")

    # Mark paragraph mark as deleted
    pPr = para._element.find(qn('w:pPr'))
    if pPr is None:
        pPr = etree.SubElement(para._element, qn('w:pPr'))
        para._element.insert(0, pPr)
    rPr = pPr.find(qn('w:rPr'))
    if rPr is None:
        rPr = etree.SubElement(pPr, qn('w:rPr'))
    del_mark = etree.SubElement(rPr, qn('w:del'))
    del_mark.set(qn('w:id'), str(os.urandom(2)[0] % 1000))
    del_mark.set(qn('w:author'), author)
    del_mark.set(qn('w:date'), date_str)

    # Wrap all runs in <w:del>
    for run_elem in list(para._element.findall(qn('w:r'))):
        del_elem = etree.Element(qn('w:del'))
        del_elem.set(qn('w:id'), str(os.urandom(2)[0] % 1000))
        del_elem.set(qn('w:author'), author)
        del_elem.set(qn('w:date'), date_str)

        para._element.replace(run_elem, del_elem)
        # Change <w:t> to <w:delText>
        for t_elem in run_elem.findall(qn('w:t')):
            t_elem.tag = qn('w:delText')
        del_elem.append(run_elem)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)

    input_path = sys.argv[1]
    output_path = sys.argv[2]
    args = sys.argv[3:]

    if not os.path.exists(input_path):
        print(f"Error: file not found: {input_path}")
        sys.exit(1)

    doc = Document(input_path)

    # Parse commands
    i = 0
    while i < len(args):
        cmd = args[i]

        if cmd == '--range':
            if i + 3 >= len(args):
                print("Error: --range requires --replace", file=sys.stderr)
                sys.exit(1)
            range_str = args[i + 1]
            # --range requires --replace
            if args[i + 2] != '--replace':
                print("Error: --range requires --replace", file=sys.stderr)
                sys.exit(1)
            old = args[i + 3]
            new = args[i + 4] if i + 4 < len(args) else ''

            # Parse and validate range N-M
            m = re.match(r'^(\d+)-(\d+)$', range_str)
            if not m:
                print(f"Error: invalid range format '{range_str}'. Expected N-M (e.g., 2-5)", file=sys.stderr)
                sys.exit(1)
            n = int(m.group(1))
            r_end = int(m.group(2))
            para_count = len(doc.paragraphs)

            if n < 0 or r_end >= para_count:
                print(f"Error: range {n}-{r_end} out of bounds (document has paragraphs 0-{para_count - 1})", file=sys.stderr)
                sys.exit(1)
            if n > r_end:
                print(f"Error: invalid range: start ({n}) > end ({r_end})", file=sys.stderr)
                sys.exit(1)

            # Range-limited replace (body paragraphs only, no tables/headers)
            count = 0
            for idx in range(n, r_end + 1):
                if replace_in_paragraph(doc.paragraphs[idx], old, new):
                    count += 1

            print(f"Range [{n}-{r_end}]: replaced '{old}' -> '{new}': {count} occurrences")
            i += 5

        elif cmd == '--replace' and i + 2 < len(args):
            old = args[i + 1]
            new = args[i + 2]
            count = replace_all(doc, old, new)
            print(f"Replace '{old}' -> '{new}': {count} occurrences")
            i += 3

        elif cmd == '--insert-after' and i + 2 < len(args):
            idx = int(args[i + 1])
            text = args[i + 2]
            insert_paragraph_after(doc, idx, text)
            i += 3

        elif cmd == '--delete' and i + 1 < len(args):
            idx = int(args[i + 1])
            delete_paragraph(doc, idx)
            i += 2

        elif cmd == '--set-style' and i + 2 < len(args):
            idx = int(args[i + 1])
            style_name = args[i + 2]
            if idx < 0 or idx >= len(doc.paragraphs):
                print(f"Error: paragraph index {idx} out of range (0-{len(doc.paragraphs)-1})", file=sys.stderr)
                sys.exit(1)
            doc.paragraphs[idx].style = doc.styles[style_name]
            print(f"Set paragraph {idx} style to '{style_name}'")
            i += 3

        elif cmd == '--list':
            max_show = int(args[i + 1]) if i + 1 < len(args) and args[i + 1].isdigit() else None
            list_paragraphs(doc, max_show)
            i += 2 if max_show else 1

        elif cmd == '--fill':
            # Read replacements from a json file: {"old1": "new1", "old2": "new2"}
            import json
            if i + 1 >= len(args):
                print("Error: --fill requires a JSON file path")
                sys.exit(1)
            json_path = args[i + 1]
            with open(json_path, 'r', encoding='utf-8') as f:
                replacements = json.load(f)
            for old, new in replacements.items():
                count = replace_all(doc, old, new)
                print(f"Replace '{old}' -> '{new}': {count} occurrences")
            i += 2

        elif cmd == '--paragraph' and i + 2 < len(args):
            idx = int(args[i + 1])
            # Look ahead for sub-options: text, --file, --append, --prepend
            j = i + 2
            text_val = None
            file_path = None
            append_val = None
            prepend_val = None

            while j < len(args):
                if args[j] == '--file' and j + 1 < len(args):
                    file_path = args[j + 1]
                    j += 2
                elif args[j] == '--append' and j + 1 < len(args):
                    append_val = args[j + 1]
                    j += 2
                elif args[j] == '--prepend' and j + 1 < len(args):
                    prepend_val = args[j + 1]
                    j += 2
                elif args[j].startswith('--'):
                    break
                else:
                    text_val = args[j]
                    j += 1

            # Validate conflicts
            if file_path and text_val:
                print("Error: --file and inline text are mutually exclusive", file=sys.stderr)
                sys.exit(1)
            if append_val and prepend_val:
                print("Error: --append and --prepend are mutually exclusive", file=sys.stderr)
                sys.exit(1)

            # Validate index
            if idx < 0 or idx >= len(doc.paragraphs):
                print(f"Error: paragraph index {idx} out of range (0-{len(doc.paragraphs)-1})", file=sys.stderr)
                sys.exit(1)

            para = doc.paragraphs[idx]

            if file_path:
                try:
                    with open(file_path, 'r', encoding='utf-8') as f:
                        new_text = f.read()
                except FileNotFoundError:
                    print(f"Error: file not found: {file_path}", file=sys.stderr)
                    sys.exit(1)
                set_paragraph_text(para, new_text)
                print(f"Set paragraph {idx} to content of '{file_path}'")
            elif append_val is not None:
                current_text = para.text
                new_text = current_text + " " + append_val
                set_paragraph_text(para, new_text)
                print(f"Appended to paragraph {idx}")
            elif prepend_val is not None:
                current_text = para.text
                new_text = prepend_val + " " + current_text
                set_paragraph_text(para, new_text)
                print(f"Prepended to paragraph {idx}")
            elif text_val is not None:
                set_paragraph_text(para, text_val)
                print(f"Set paragraph {idx} to '{text_val}'")

            i = j

        else:
            print(f"Unknown command: {cmd}")
            i += 1

    doc.save(output_path)
    print(f"\nSaved: {output_path}")


if __name__ == '__main__':
    main()
