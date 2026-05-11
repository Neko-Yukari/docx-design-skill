#!/usr/bin/env python3
"""
Extract text from .docx files into structured Markdown.
Usage: python extract_docx.py <input.docx> [output.md]
       python extract_docx.py <input.docx> --structure [--detect-freeform]
       python extract_docx.py <input.docx> --range N-M [output.md]
       python extract_docx.py <input.docx> --section "name" [--exact] [output.md]
       python extract_docx.py <input.docx> --grep "pattern" [--context N]

If output is omitted, writes to <input>_extracted.md (or <input>_range_N_M.md,
<input>_section_<name>.md for --range/--section)
"""
import sys
import os
import re
from docx import Document

HEADING_PATTERNS = [
    (r'^第[一二三四五六七八九十\d]+章', 1, 0.95),
    (r'^第[一二三四五六七八九十\d]+节', 2, 0.95),
    (r'^\d+\.\d+\s+\w', 2, 0.9),
    (r'^\d+\.\s+\w', 1, 0.85),
    (r'^\d+\s+\w', 1, 0.8),
    (r'^[一二三四五六七八九十]、', 1, 0.8),
    (r'^(Abstract|摘要|Introduction|引言|Conclusion|结论|References|参考文献|Acknowledgements|致谢)$', 1, 0.75),
]


def extract_docx_to_markdown(docx_path, output_path):
    doc = Document(docx_path)

    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(f"# Source: {os.path.basename(docx_path)}\n\n")

        for para in doc.paragraphs:
            style = para.style.name if para.style else 'Normal'
            text = para.text.strip()
            if not text:
                f.write('\n')
                continue

            if style.startswith('Heading'):
                try:
                    level = int(style.replace('Heading ', ''))
                    f.write(f"{'#' * level} {text}\n\n")
                except ValueError:
                    f.write(f"# {text}\n\n")
            else:
                f.write(f"{text}\n\n")

        # Tables
        if doc.tables:
            f.write(f"\n---\n\n")
        for i, table in enumerate(doc.tables):
            f.write(f"\n**Table {i + 1}:**\n\n")
            for row in table.rows:
                cells = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
                f.write('| ' + ' | '.join(cells) + ' |\n')
            f.write('\n')

    print(f"Extracted: {docx_path} -> {output_path}")
    print(f"  Paragraphs: {len(doc.paragraphs)}")
    print(f"  Tables: {len(doc.tables)}")


def _get_body_order(doc):
    """Return list of ('para', idx) or ('table', idx) in document order."""
    order = []
    para_idx = 0
    table_idx = 0
    for child in doc.element.body:
        tag = child.tag
        if tag.endswith('}p'):
            order.append(('para', para_idx))
            para_idx += 1
        elif tag.endswith('}tbl'):
            order.append(('table', table_idx))
            table_idx += 1
    return order


def extract_structure(doc, detect_freeform=False):
    """
    Extract document heading structure.

    Returns a list of (para_index, heading_level, text, confidence, reason) tuples.
    Each tuple represents a detected heading in document order.
    """
    structures = []
    style_detected = set()

    # Style-based detection (always runs)
    for i, para in enumerate(doc.paragraphs):
        style = para.style.name if para.style else ''
        if style.startswith('Heading'):
            try:
                level = int(style.replace('Heading ', ''))
            except ValueError:
                level = 1
            text = para.text.strip()
            if text:
                structures.append((i, level, text, 1.0, f'style:{style}'))
                style_detected.add(i)

    # Free-form detection via regex patterns
    if detect_freeform:
        for i, para in enumerate(doc.paragraphs):
            if i in style_detected:
                continue
            text = para.text.strip()
            if not text:
                continue

            for pattern, level, confidence in HEADING_PATTERNS:
                if re.match(pattern, text):
                    # Context validation for low-confidence patterns
                    if confidence < 0.9:
                        # Need next paragraph to exist and be longer
                        if i + 1 >= len(doc.paragraphs):
                            continue
                        next_text = doc.paragraphs[i + 1].text.strip()
                        if len(next_text) <= len(text):
                            continue
                        # Word count check
                        if len(text.split()) > 10:
                            continue

                    structures.append((i, level, text, confidence, f'pattern:{pattern}'))
                    break  # First matching pattern wins

    # Sort by paragraph index
    structures.sort(key=lambda x: x[0])
    return structures


def print_structure(structure, doc, filename="document.docx"):
    """Print heading tree to stdout."""
    para_count = len(doc.paragraphs)
    table_count = len(doc.tables)

    print(f"[document] {filename} ({para_count} paragraphs, {table_count} tables)")

    if not structure:
        return

    # Build body element order for table-position lookups
    body_order = _get_body_order(doc)

    para_body_pos = {}
    table_body_pos = {}
    for pos, (etype, eidx) in enumerate(body_order):
        if etype == 'para':
            para_body_pos[eidx] = pos
        else:
            table_body_pos[eidx] = pos

    n = len(structure)
    for idx, (pi, level, text, confidence, reason) in enumerate(structure):
        # End paragraph of this heading's range
        if idx + 1 < n:
            end_pi = structure[idx + 1][0] - 1
        else:
            end_pi = len(doc.paragraphs) - 1

        # Find tables between this heading and next heading in body order
        heading_body_pos = para_body_pos.get(pi, -1)
        if idx + 1 < n:
            next_pi = structure[idx + 1][0]
            next_body_pos = para_body_pos.get(next_pi, len(body_order))
        else:
            next_body_pos = len(body_order)

        tables_in_range = []
        for ti, tpos in table_body_pos.items():
            if heading_body_pos < tpos < next_body_pos:
                tables_in_range.append(ti)

        # Build range string
        range_parts = [f"p{pi}-p{end_pi}"]
        for ti in tables_in_range:
            range_parts.append(f"table:{ti}")
        range_str = ", ".join(range_parts)

        # Build indent (based on heading level)
        indent = "|   " * (level - 1)

        # Build heading label and optional confidence
        is_freeform = confidence < 1.0
        if is_freeform:
            heading_label = f"H{level}"
            conf_str = f" [conf:{confidence}]"
        else:
            heading_label = f"Heading {level}"
            conf_str = ""

        line = f"{indent}+-- {heading_label} (p{pi}){conf_str}: {text} [{range_str}]"
        print(line)


def _truncate(text, max_len):
    """Truncate text to max_len chars, appending '...' if truncated."""
    if len(text) > max_len:
        return text[:max_len] + '...'
    return text


def grep_paragraphs(doc, pattern, context=0):
    """Search paragraphs for pattern, output matches with optional context to stdout."""
    paragraphs = doc.paragraphs
    matches = []

    for i, para in enumerate(paragraphs):
        if re.search(pattern, para.text, re.IGNORECASE):
            matches.append(i)

    if not matches:
        print(f"No matches found for '{pattern}'")
        return

    for idx, pi in enumerate(matches):
        if idx > 0:
            print()

        # Matched paragraph
        text = paragraphs[pi].text.strip()
        display = _truncate(text, 120)
        print(f'[paragraph {pi}] "{display}"')

        # Context before (closest first)
        for offset in range(1, context + 1):
            prev_pi = pi - offset
            if prev_pi >= 0:
                prev_text = paragraphs[prev_pi].text.strip()
                prev_display = _truncate(prev_text, 120)
                print(f'  <- p{prev_pi}: "{prev_display}"')

        # Context after (closest first)
        for offset in range(1, context + 1):
            next_pi = pi + offset
            if next_pi < len(paragraphs):
                next_text = paragraphs[next_pi].text.strip()
                next_display = _truncate(next_text, 120)
                print(f'  -> p{next_pi}: "{next_display}"')


def main():
    if len(sys.argv) < 2:
        print("Usage: python extract_docx.py <input.docx> [output.md]")
        print("       python extract_docx.py <input.docx> --structure [--detect-freeform]")
        print("       python extract_docx.py <input.docx> --range N-M [output.md]")
        print("       python extract_docx.py <input.docx> --section \"name\" [--exact] [output.md]")
        print("       python extract_docx.py <input.docx> --grep \"pattern\" [--context N]")
        sys.exit(1)

    input_path = sys.argv[1]
    args = sys.argv[2:]

    use_structure = '--structure' in args
    detect_freeform = '--detect-freeform' in args

    if detect_freeform and not use_structure:
        print("Error: --detect-freeform requires --structure")
        sys.exit(1)

    if use_structure:
        doc = Document(input_path)
        structure = extract_structure(doc, detect_freeform=detect_freeform)
        filename = os.path.basename(input_path)
        print_structure(structure, doc, filename=filename)
        sys.exit(0)

    use_grep = '--grep' in args
    if use_grep:
        grep_idx = args.index('--grep')
        if grep_idx + 1 >= len(args):
            sys.stderr.write("Error: --grep requires a pattern (e.g., --grep \"keyword\")\n")
            sys.exit(1)
        pattern = args[grep_idx + 1]

        context = 0
        if '--context' in args:
            ctx_idx = args.index('--context')
            if ctx_idx + 1 >= len(args):
                sys.stderr.write("Error: --context requires a number\n")
                sys.exit(1)
            try:
                context = int(args[ctx_idx + 1])
            except ValueError:
                sys.stderr.write(f"Error: --context value must be an integer, got '{args[ctx_idx + 1]}'\n")
                sys.exit(1)
            if context < 0:
                sys.stderr.write(f"Error: --context must be >= 0, got {context}\n")
                sys.exit(1)

        doc = Document(input_path)
        grep_paragraphs(doc, pattern, context)
        sys.exit(0)

    use_range = '--range' in args
    if use_range:
        range_idx = args.index('--range')
        if range_idx + 1 >= len(args):
            sys.stderr.write("Error: --range requires a value (e.g., --range 2-5)\n")
            sys.exit(1)
        range_str = args[range_idx + 1]

        match = re.match(r'^(\d+)-(\d+)$', range_str)
        if not match:
            sys.stderr.write(f"Error: invalid range format '{range_str}'. Use N-M (e.g., 2-5)\n")
            sys.exit(1)
        n = int(match.group(1))
        m = int(match.group(2))

        if n < 1:
            sys.stderr.write(f"Error: range start must be >= 1, got {n}\n")
            sys.exit(1)
        if n > m:
            sys.stderr.write(f"Error: invalid range: start ({n}) > end ({m})\n")
            sys.exit(1)

        doc = Document(input_path)
        total_paras = len(doc.paragraphs)
        if m > total_paras:
            sys.stderr.write(f"Error: range end ({m}) exceeds document paragraph count ({total_paras})\n")
            sys.exit(1)

        # Determine output path: explicit positional arg (skipping the range value itself)
        positional = [a for a in args if not a.startswith('--') and a != range_str]
        if positional:
            output_path = positional[0]
        else:
            base = os.path.splitext(input_path)[0]
            output_path = f"{base}_range_{n}_{m}.md"

        # Write range extraction
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(f"# Source: {os.path.basename(input_path)} (paragraphs {n}-{m})\n\n")

            for para in doc.paragraphs[n - 1:m]:
                style = para.style.name if para.style else 'Normal'
                text = para.text.strip()
                if not text:
                    f.write('\n')
                    continue

                if style.startswith('Heading'):
                    try:
                        level = int(style.replace('Heading ', ''))
                        f.write(f"{'#' * level} {text}\n\n")
                    except ValueError:
                        f.write(f"# {text}\n\n")
                else:
                    f.write(f"{text}\n\n")

            # Tables between range paragraphs in document body order
            body_order = _get_body_order(doc)
            first_body_pos = None
            last_body_pos = None
            for pos, (etype, eidx) in enumerate(body_order):
                if etype == 'para' and n - 1 <= eidx < m:
                    if first_body_pos is None:
                        first_body_pos = pos
                    last_body_pos = pos

            tables_in_range = []
            if first_body_pos is not None and last_body_pos is not None:
                for pos, (etype, eidx) in enumerate(body_order):
                    if etype == 'table' and first_body_pos < pos < last_body_pos:
                        tables_in_range.append(eidx)

            if tables_in_range:
                f.write(f"\n---\n\n")
            for ti in tables_in_range:
                table = doc.tables[ti]
                f.write(f"\n**Table {ti + 1}:**\n\n")
                for row in table.rows:
                    cells = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
                    f.write('| ' + ' | '.join(cells) + ' |\n')
                f.write('\n')

        print(f"Extracted: {input_path}[{n}-{m}] -> {output_path}")
        sys.exit(0)

    use_section = '--section' in args
    if use_section:
        section_idx = args.index('--section')
        if section_idx + 1 >= len(args):
            sys.stderr.write("Error: --section requires a value (e.g., --section \"Introduction\")\n")
            sys.exit(1)
        section_name = args[section_idx + 1]
        use_exact = '--exact' in args

        doc = Document(input_path)
        structure = extract_structure(doc)

        # Find matching headings
        matched = []
        for pi, level, text, _confidence, _reason in structure:
            if use_exact:
                if text == section_name:
                    matched.append((pi, level, text))
            else:
                if section_name.lower() in text.lower():
                    matched.append((pi, level, text))

        if not matched:
            sys.stderr.write(f"Error: no section matching '{section_name}'\n")
            if structure:
                sys.stderr.write("Available headings:\n")
                for _pi, _level, text, _confidence, _reason in structure:
                    sys.stderr.write(f"  H{_level}: {text}\n")
            sys.exit(1)

        # Determine output path
        positional = [a for a in args if not a.startswith('--') and a != section_name]
        if positional:
            output_path = positional[0]
        else:
            safe_name = re.sub(r'[^\w]', '_', section_name)[:30]
            base = os.path.splitext(input_path)[0]
            output_path = f"{base}_section_{safe_name}.md"

        # Extract all matching sections
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(f"# Source: {os.path.basename(input_path)} (section: {section_name})\n\n")

            for match_idx, (start_pi, match_level, match_text) in enumerate(matched):
                if match_idx > 0:
                    f.write("\n---\n\n")

                # Find end paragraph: next heading with level <= match_level, or end of document
                end_pi = len(doc.paragraphs) - 1
                for pi, level, _text, _confidence, _reason in structure:
                    if pi > start_pi and level <= match_level:
                        end_pi = pi - 1
                        break

                # Write paragraphs in the section range
                for pi in range(start_pi, end_pi + 1):
                    para = doc.paragraphs[pi]
                    style = para.style.name if para.style else 'Normal'
                    text = para.text.strip()
                    if not text:
                        f.write('\n')
                        continue

                    if style.startswith('Heading'):
                        try:
                            level = int(style.replace('Heading ', ''))
                            f.write(f"{'#' * level} {text}\n\n")
                        except ValueError:
                            f.write(f"# {text}\n\n")
                    else:
                        f.write(f"{text}\n\n")

                # Tables within the section range (document body order)
                body_order = _get_body_order(doc)
                para_body_pos = {}
                table_body_pos = {}
                for pos, (etype, eidx) in enumerate(body_order):
                    if etype == 'para':
                        para_body_pos[eidx] = pos
                    else:
                        table_body_pos[eidx] = pos

                first_body_pos = None
                last_body_pos = None
                for pos, (etype, eidx) in enumerate(body_order):
                    if etype == 'para' and start_pi <= eidx <= end_pi:
                        if first_body_pos is None:
                            first_body_pos = pos
                        last_body_pos = pos

                tables_in_section = []
                if first_body_pos is not None and last_body_pos is not None:
                    for pos, (etype, eidx) in enumerate(body_order):
                        if etype == 'table' and first_body_pos < pos < last_body_pos:
                            tables_in_section.append(eidx)

                if tables_in_section:
                    f.write(f"\n---\n\n")
                for ti in tables_in_section:
                    table = doc.tables[ti]
                    f.write(f"\n**Table {ti + 1}:**\n\n")
                    for row in table.rows:
                        cells = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
                        f.write('| ' + ' | '.join(cells) + ' |\n')
                    f.write('\n')

        print(f"Extracted: {input_path}[section='{section_name}'] -> {output_path}")
        print(f"  Matched {len(matched)} heading(s)")
        sys.exit(0)

    # Original behavior: remaining positional args are output path
    positional = [a for a in args if not a.startswith('--')]
    if positional:
        output_path = positional[0]
    else:
        base = os.path.splitext(input_path)[0]
        output_path = f"{base}_extracted.md"

    extract_docx_to_markdown(input_path, output_path)


if __name__ == '__main__':
    main()
