#!/usr/bin/env python3
"""
Convert LaTeX formulas to OMML for .docx embedding.
Uses latex2mathml for robust LaTeX parsing, then converts MathML to OMML.

Usage:
    from latex2omml import latex_to_omml, insert_formula

    omml = latex_to_omml(r"x = \frac{-b \pm \sqrt{b^2 - 4ac}}{2a}")

    from docx import Document
    doc = Document()
    p = doc.add_paragraph()
    insert_formula(p, r"E = mc^2")
    doc.save('output.docx')

CLI: python latex2omml.py "E = mc^2"
"""
import sys
from xml.etree.ElementTree import Element, SubElement, tostring, fromstring
from latex2mathml import converter as _latex_converter

OMML_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
MATHML_NS = 'http://www.w3.org/1998/Math/MathML'
XML_NS = 'http://www.w3.org/XML/1998/namespace'


def _tag(elem):
    return elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag


def _extract_text(elem):
    parts = []
    if elem.text: parts.append(elem.text)
    for c in elem:
        parts.append(_extract_text(c))
        if c.tail: parts.append(c.tail)
    return ''.join(parts)


def _add_text_run(parent, text):
    r = SubElement(parent, f'{{{OMML_NS}}}r')
    t = SubElement(r, f'{{{OMML_NS}}}t')
    t.text = text
    if text.startswith(' ') or text.endswith(' '):
        t.set(f'{{{XML_NS}}}space', 'preserve')


def _convert(ml_elem, omml_parent):
    tag = _tag(ml_elem)
    kids = list(ml_elem)

    if tag == 'mrow':
        for c in kids: _convert(c, omml_parent)

    elif tag == 'msup':
        e = SubElement(omml_parent, f'{{{OMML_NS}}}sSup')
        SubElement(e, f'{{{OMML_NS}}}sSupPr')
        ce = SubElement(e, f'{{{OMML_NS}}}e'); cs = SubElement(e, f'{{{OMML_NS}}}sup')
        if kids: _convert(kids[0], ce)
        if len(kids) > 1: _convert(kids[1], cs)

    elif tag == 'msub':
        e = SubElement(omml_parent, f'{{{OMML_NS}}}sSub')
        SubElement(e, f'{{{OMML_NS}}}sSubPr')
        ce = SubElement(e, f'{{{OMML_NS}}}e'); cs = SubElement(e, f'{{{OMML_NS}}}sub')
        if kids: _convert(kids[0], ce)
        if len(kids) > 1: _convert(kids[1], cs)

    elif tag == 'msubsup':
        e = SubElement(omml_parent, f'{{{OMML_NS}}}sSubSup')
        SubElement(e, f'{{{OMML_NS}}}sSubSupPr')
        ce = SubElement(e, f'{{{OMML_NS}}}e')
        csub = SubElement(e, f'{{{OMML_NS}}}sub')
        csup = SubElement(e, f'{{{OMML_NS}}}sup')
        if kids: _convert(kids[0], ce)
        if len(kids) > 1: _convert(kids[1], csub)
        if len(kids) > 2: _convert(kids[2], csup)

    elif tag == 'mfrac':
        e = SubElement(omml_parent, f'{{{OMML_NS}}}f')
        SubElement(e, f'{{{OMML_NS}}}fPr')
        num = SubElement(e, f'{{{OMML_NS}}}num')
        den = SubElement(e, f'{{{OMML_NS}}}den')
        if kids: _convert(kids[0], num)
        if len(kids) > 1: _convert(kids[1], den)

    elif tag == 'msqrt':
        e = SubElement(omml_parent, f'{{{OMML_NS}}}rad')
        SubElement(e, f'{{{OMML_NS}}}radPr')
        SubElement(e, f'{{{OMML_NS}}}deg')
        ce = SubElement(e, f'{{{OMML_NS}}}e')
        for c in kids: _convert(c, ce)

    elif tag == 'mroot':
        e = SubElement(omml_parent, f'{{{OMML_NS}}}rad')
        SubElement(e, f'{{{OMML_NS}}}radPr')
        deg = SubElement(e, f'{{{OMML_NS}}}deg')
        ce = SubElement(e, f'{{{OMML_NS}}}e')
        if len(kids) > 1: _convert(kids[1], deg)
        if kids: _convert(kids[0], ce)

    elif tag == 'munderover':
        op = _extract_text(kids[0]) if kids else ''
        nary_ops = {'\u222B': '\u222B', '\u2211': '\u2211', '\u220F': '\u220F', '\u222E': '\u222E'}
        if op in nary_ops:
            e = SubElement(omml_parent, f'{{{OMML_NS}}}nary')
            pr = SubElement(e, f'{{{OMML_NS}}}naryPr')
            SubElement(pr, f'{{{OMML_NS}}}chr').set(f'{{{OMML_NS}}}val', nary_ops[op])
            se = SubElement(e, f'{{{OMML_NS}}}sub'); pe = SubElement(e, f'{{{OMML_NS}}}sup'); ee = SubElement(e, f'{{{OMML_NS}}}e')
            if len(kids) > 1: _convert(kids[1], se)
            if len(kids) > 2: _convert(kids[2], pe)
            if len(kids) > 3: _convert(kids[3], ee)
        else:
            e = SubElement(omml_parent, f'{{{OMML_NS}}}sSubSup')
            SubElement(e, f'{{{OMML_NS}}}sSubSupPr')
            ce = SubElement(e, f'{{{OMML_NS}}}e'); csub = SubElement(e, f'{{{OMML_NS}}}sub'); csup = SubElement(e, f'{{{OMML_NS}}}sup')
            if kids: _convert(kids[0], ce)
            if len(kids) > 1: _convert(kids[1], csub)
            if len(kids) > 2: _convert(kids[2], csup)

    elif tag == 'munder':
        e = SubElement(omml_parent, f'{{{OMML_NS}}}sSub')
        SubElement(e, f'{{{OMML_NS}}}sSubPr')
        ce = SubElement(e, f'{{{OMML_NS}}}e'); cs = SubElement(e, f'{{{OMML_NS}}}sub')
        if kids: _convert(kids[0], ce)
        if len(kids) > 1: _convert(kids[1], cs)

    elif tag == 'mover':
        e = SubElement(omml_parent, f'{{{OMML_NS}}}sSup')
        SubElement(e, f'{{{OMML_NS}}}sSupPr')
        ce = SubElement(e, f'{{{OMML_NS}}}e'); cs = SubElement(e, f'{{{OMML_NS}}}sup')
        if kids: _convert(kids[0], ce)
        if len(kids) > 1: _convert(kids[1], cs)

    elif tag in ('mi', 'mn', 'mo', 'mtext'):
        t = _extract_text(ml_elem)
        if t: _add_text_run(omml_parent, t)

    elif tag == 'mspace':
        _add_text_run(omml_parent, ' ')

    elif tag == 'mfenced':
        op = ml_elem.get('open', '('); cl = ml_elem.get('close', ')')
        _add_text_run(omml_parent, op)
        for c in kids: _convert(c, omml_parent)
        _add_text_run(omml_parent, cl)

    elif tag == 'none':
        pass

    else:
        for c in kids: _convert(c, omml_parent)


def mathml_to_omml(mathml_root):
    root = Element(f'{{{OMML_NS}}}oMathPara')
    for child in list(mathml_root):
        t = _tag(child)
        if t in ('mrow', 'mstyle'):
            om = SubElement(root, f'{{{OMML_NS}}}oMath')
            _convert(child, om)
        elif t == 'semantics':
            for sc in child:
                if _tag(sc) == 'mrow':
                    om = SubElement(root, f'{{{OMML_NS}}}oMath')
                    _convert(sc, om)
                    break
    return root


def latex_to_omml(latex_str):
    ml = _latex_converter.convert(latex_str)
    if 'xmlns' not in ml[:200]:
        ml = ml.replace('<math', f'<math xmlns="{MATHML_NS}"', 1)
    return tostring(mathml_to_omml(fromstring(ml)), encoding='unicode')


def insert_formula(paragraph, latex_str):
    from lxml import etree
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run('')
    run._element.append(etree.fromstring(latex_to_omml(latex_str)))
    return run


if __name__ == '__main__':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

    if len(sys.argv) < 2:
        print("Usage: python latex2omml.py 'E = mc^2'")
        print("       python latex2omml.py --file formulas.txt")
        sys.exit(1)
    if sys.argv[1] == '--file':
        with open(sys.argv[2], 'r', encoding='utf-8') as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith('#'):
                    print("--- " + line)
                    print(latex_to_omml(line))
                    print()
    else:
        print(latex_to_omml(' '.join(sys.argv[1:])))
